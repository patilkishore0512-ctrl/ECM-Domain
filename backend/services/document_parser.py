import os
import pdfplumber
import openpyxl
from typing import List, Dict

ASSET_KEYWORDS = {
    "transformer": [
        "transformer", "winding", "buchholz", "tap changer", "oltc", "dga",
        "dissolved gas", "oil temperature", "core", "onan", "onaf", "onan/onaf",
        "turns ratio", "magnetizing", "inrush", "differential protection", "REF",
        "oil level", "conservator", "silica gel", "bushing", "vector group"
    ],
    "switchgear": [
        "switchgear", "circuit breaker", "sf6", "busbar", "arc flash",
        "contactor", "switchboard", "MV panel", "LV panel", "ACB", "VCB",
        "MCB", "MCCB", "partial discharge", "operating mechanism", "CB",
        "earthing switch", "isolator", "withdrawable", "fixed pattern"
    ],
    "ups": [
        "ups", "uninterruptible", "battery", "inverter", "rectifier",
        "charger", "igbt", "bypass", "static switch", "float charge",
        "equalize charge", "SOC", "state of charge", "VRLA", "AGM",
        "back feed", "load bank", "autonomy time", "kva", "input filter"
    ],
    "motors": [
        "motor", "rotor", "stator", "bearing", "shaft", "insulation resistance",
        "slip", "induction", "synchronous", "winding temperature", "vibration",
        "alignment", "coupling", "phase unbalance", "locked rotor",
        "overload relay", "thermistor", "PTC", "NTC", "megger", "polarization index"
    ]
}

CHUNK_SIZE = 800
CHUNK_OVERLAP = 150


def _classify_chunk(text: str) -> str:
    text_lower = text.lower()
    scores = {asset: 0 for asset in ASSET_KEYWORDS}
    for asset, keywords in ASSET_KEYWORDS.items():
        for kw in keywords:
            if kw.lower() in text_lower:
                scores[asset] += 1
    best = max(scores, key=scores.get)
    if scores[best] == 0:
        return "general"
    return best


def _extract_tables_from_page(page) -> str:
    tables = page.extract_tables()
    if not tables:
        return ""
    result = []
    for table in tables:
        for row in table:
            cleaned = [str(cell).strip() if cell else "" for cell in row]
            if any(c for c in cleaned):
                result.append(" | ".join(cleaned))
    return "\n".join(result)


def parse_knowledge_pdf(pdf_path: str) -> List[Dict]:
    chunks = []
    raw_pages = []

    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            tables = _extract_tables_from_page(page)
            combined = f"{text}\n{tables}".strip()
            if combined:
                raw_pages.append({"page": page_num, "text": combined})

    # Chunk pages with overlap
    for page_data in raw_pages:
        text = page_data["text"]
        words = text.split()
        start = 0
        while start < len(words):
            chunk_words = words[start: start + CHUNK_SIZE]
            chunk_text = " ".join(chunk_words)
            asset_type = _classify_chunk(chunk_text)
            chunks.append({
                "text": chunk_text,
                "asset_type": asset_type,
                "page": page_data["page"],
                "source": os.path.basename(pdf_path)
            })
            start += CHUNK_SIZE - CHUNK_OVERLAP

    return chunks


def parse_tech_spec(file_path: str, file_type: str) -> str:
    if file_type == "pdf":
        return _parse_spec_pdf(file_path)
    elif file_type == "xlsx":
        return _parse_spec_excel(file_path)
    elif file_type == "docx":
        return _parse_spec_word(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def _parse_spec_pdf(file_path: str) -> str:
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            tables = _extract_tables_from_page(page)
            combined = f"{text}\n{tables}".strip()
            if combined:
                text_parts.append(combined)
    return "\n\n".join(text_parts)


def _parse_spec_excel(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text_parts = []
    for sheet in wb.worksheets:
        text_parts.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(c for c in cells):
                text_parts.append(" | ".join(cells))
    return "\n".join(text_parts)


def _parse_spec_word(file_path: str) -> str:
    try:
        import docx
        doc = docx.Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        raise ImportError("python-docx not installed. Run: pip install python-docx")
